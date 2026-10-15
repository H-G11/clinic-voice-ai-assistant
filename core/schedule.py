"""Шаблон Excel. НЕ подменяйте им реальное расписание клиники."""
from __future__ import annotations

from datetime import date, datetime, time, timedelta
from pathlib import Path
from openpyxl import Workbook, load_workbook
from .store import Store
import hashlib
import json
from openpyxl.utils.datetime import from_excel

DOCTORS = [
    ("d1", "Иванов Иван Иванович", "терапевт"),
    ("d2", "Петрова Анна Сергеевна", "невролог"),
    ("d3", "Сидоров Павел Олегович", "хирург"),
]


def write_demo_excel(path: str | Path) -> None:
    p = Path(path)
    p.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "schedule"
    ws.append(["doctor_id", "doctor_name", "specialty", "date", "time"])
    for offset in range(21):
        day = date.today() + timedelta(days=offset)
        if day.weekday() >= 5:
            continue
        for did, name, specialty in DOCTORS:
            for hour in range(9, 17):
                for minute in (0, 30):
                    ws.append([did, name, specialty, day.strftime("%Y-%m-%d"), f"{hour:02d}:{minute:02d}"])
    wb.save(p)


def import_excel(store: Store, path: str | Path) -> int:
    """Импорт только шаблона с конкретными пятью колонками.

    Не стирает существующие записи. При изменении расписания используйте
    отдельную процедуру сверки; во время демо импортируйте один раз.
    """
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        cols = [str(x).strip() if x is not None else "" for x in next(rows)]
        expected = ["doctor_id", "doctor_name", "specialty", "date", "time"]
        if cols[:5] != expected:
            raise ValueError(f"Ожидаются колонки {expected}, получены {cols[:5]}")
        doctors, slots = {}, set()
        for n, row in enumerate(rows, start=2):
            if not row or not row[0]:
                continue
            did, name, specialty, d, t = row[:5]
            did, name, specialty = str(did).strip(), str(name).strip(), str(specialty).strip().lower()
            if isinstance(d, (datetime, date)):
                day = d.strftime("%Y-%m-%d")
            else:
                day = date.fromisoformat(str(d).strip()).isoformat()
            if isinstance(t, (datetime, time)):
                tm = t.strftime("%H:%M")
            else:
                tm = datetime.strptime(str(t).strip(), "%H:%M").strftime("%H:%M")
            doctors[did] = (name, specialty)
            slots.add((did, f"{day}T{tm}"))
        if not slots:
            raise ValueError("Excel не содержит слотов")
        with store.connect() as con:
            con.execute("BEGIN IMMEDIATE")
            # Входное расписание должно быть стабильным после создания записей.
            if con.execute("SELECT COUNT(*) FROM bookings").fetchone()[0] > 0:
                raise RuntimeError("Есть брони; замена расписания запрещена без сверки")
            con.execute("DELETE FROM slots")
            con.execute("DELETE FROM doctors")
            con.executemany("INSERT INTO doctors(id,name,specialty) VALUES (?,?,?)", [(d,*v) for d,v in doctors.items()])
            con.executemany("INSERT INTO slots(doctor_id,start_at) VALUES (?,?)", sorted(slots))
            con.commit()
        return len(slots)
    finally:
        wb.close()


def excel_day(value, epoch):
    if isinstance(value, (int, float)):
        value = from_excel(value, epoch)
    if isinstance(value, (datetime, date)):
        return value.strftime('%Y-%m-%d')
    return date.fromisoformat(str(value).strip()).isoformat()


def excel_time(value, epoch):
    if isinstance(value, (int, float)):
        value = from_excel(value, epoch)
    if isinstance(value, (datetime, time)):
        return value.strftime('%H:%M')
    return datetime.strptime(str(value).strip(), '%H:%M').strftime('%H:%M')


def import_clinic_excel(store: Store, path: str | Path) -> dict:
    """Stream the supplied seven-sheet case workbook into an EMPTY separate DB.

    Preserve every source slot/status. Availability is decided at query time.
    No generated dates, no deletion of existing reservations, no implicit reimport.
    """
    path = Path(path)
    with path.open('rb') as file:
        digest = hashlib.file_digest(file,'sha256').hexdigest()
    with store.connect() as c:
        previous = c.execute("SELECT value FROM metadata WHERE key='source_sha256'").fetchone()
        if previous and previous[0] == digest:
            if not c.execute('SELECT 1 FROM source_reference LIMIT 1').fetchone():
                # Additive migration of reference provenance; never replace doctors/slots/bookings.
                archive=load_workbook(path,read_only=True,data_only=True)
                try:
                    c.execute('BEGIN IMMEDIATE')
                    _archive_reference(archive,c)
                    c.commit()
                finally: archive.close()
            return {'doctors': c.execute('SELECT COUNT(*) FROM doctors').fetchone()[0],
                    'slots': c.execute('SELECT COUNT(*) FROM slots').fetchone()[0], 'cached': True}
        if c.execute('SELECT COUNT(*) FROM doctors').fetchone()[0] or c.execute('SELECT COUNT(*) FROM bookings').fetchone()[0]:
            raise RuntimeError('Импорт разрешён только в отдельную пустую БД; нужна сверка')
    wb = load_workbook(path, read_only=True, data_only=True)
    required = {'Описание', 'Специальности', 'Врачи', 'Слоты', 'Недоступность', 'Календарь', 'График'}
    try:
        if not required.issubset(wb.sheetnames):
            raise ValueError('Нужны все семь листов расписания кейса')
        with store.connect() as c:
            c.execute('BEGIN IMMEDIATE')
            if c.execute('SELECT COUNT(*) FROM doctors').fetchone()[0]:
                raise RuntimeError('Расписание уже импортировано другим процессом')
            for ws in wb:
                header = next(ws.iter_rows(values_only=True))
                c.execute('INSERT INTO source_sheets VALUES (?,?,?)', (ws.title, ws.max_row-1, json.dumps(header, ensure_ascii=False, default=str)))
            _archive_reference(wb,c)
            # Reference/summary sheets are retained locally by source hash, not treated as availability.
            for row in wb['Врачи'].iter_rows(min_row=2, values_only=True):
                did, name, dept, specialty, patients, room, duration = row[:7]
                c.execute('INSERT INTO doctors(id,name,specialty,department,patients,room,duration) VALUES (?,?,?,?,?,?,?)',
                          (str(did), str(name).strip(), str(specialty).strip().lower(), str(dept), str(patients), str(room), int(duration)))
            for row in wb['Недоступность'].iter_rows(min_row=2, values_only=True):
                c.execute('INSERT INTO unavailability VALUES (?,?,?,?)',
                          (row[0], excel_day(row[4], wb.epoch), excel_day(row[5], wb.epoch), row[3]))
            calendar = wb['Календарь'].iter_rows(values_only=True)
            dates = next(calendar)[3:]
            next(calendar)  # weekday display row, not a doctor's calendar
            for row in calendar:
                if row[0]:
                    c.executemany('INSERT INTO calendar VALUES (?,?,?)',
                                  [(row[0], excel_day(d, wb.epoch), code) for d,code in zip(dates,row[3:]) if d is not None])
            headers = next(wb['Слоты'].iter_rows(values_only=True))
            expected = ('ID слота','Дата','День','ID врача','ФИО врача','Специальность','Отделение','Кабинет','Начало','Окончание','Мин','Статус')
            if tuple(headers) != expected:
                raise ValueError('Изменилась схема листа Слоты')
            batch, count = [], 0
            for row in wb['Слоты'].iter_rows(min_row=2, values_only=True):
                sid, d, _, did, name, spec, dept, room, start, end, duration, status = row
                day = excel_day(d, wb.epoch)
                start_at, end_at = f'{day}T{excel_time(start, wb.epoch)}', f'{day}T{excel_time(end, wb.epoch)}'
                if int(duration) <= 0 or datetime.fromisoformat(end_at)-datetime.fromisoformat(start_at) != timedelta(minutes=int(duration)):
                    raise ValueError(f'Некорректная длительность слота {sid}')
                batch.append((did,start_at,sid,end_at,int(duration),str(room),str(status)))
                count += 1
                if len(batch) >= 2000:
                    c.executemany('INSERT INTO slots(doctor_id,start_at,source_id,end_at,duration,room,file_status) VALUES (?,?,?,?,?,?,?)', batch)
                    batch.clear()
            if batch:
                c.executemany('INSERT INTO slots(doctor_id,start_at,source_id,end_at,duration,room,file_status) VALUES (?,?,?,?,?,?,?)', batch)
            c.executemany('INSERT INTO metadata VALUES (?,?)', [('dataset','clinic_excel'), ('source_sha256',digest), ('timezone','GMT+4')])
            c.commit()
            return {'doctors':c.execute('SELECT COUNT(*) FROM doctors').fetchone()[0], 'slots':count, 'cached':False}
    finally:
        wb.close()


def _archive_reference(wb,con):
    for ws in wb:
        if ws.title=='Слоты': continue  # slots already stored as indexed typed rows
        con.executemany('INSERT OR IGNORE INTO source_reference VALUES (?,?,?)',
                        ((ws.title,n,json.dumps(row,ensure_ascii=False,default=str)) for n,row in enumerate(ws.iter_rows(values_only=True),1)))
